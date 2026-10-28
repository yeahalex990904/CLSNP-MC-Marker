"""Build coloured results workbook (F6 UT2 style, openpyxl)."""
from __future__ import annotations

from collections import defaultdict
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL

FILL_HEADER = PatternFill("solid", fgColor="1F4E78")
FILL_KEY = PatternFill("solid", fgColor="FFF2CC")
FILL_SUM = PatternFill("solid", fgColor="DDEBF7")
FILL_GREEN = PatternFill("solid", fgColor="C6EFCE")
FILL_RED = PatternFill("solid", fgColor="FFC7CE")
FILL_YELLOW = PatternFill("solid", fgColor="FFEB9C")

COLOUR_MAP = {"green": FILL_GREEN, "red": FILL_RED, "yellow": FILL_YELLOW}


def build_results_xlsx(
    students: list[dict[str, Any]],
    key: list[str],
    title: str,
    out_path: str,
) -> dict:
    """
    students: list of {
      class_code, class_no, name, answers: [{ans, status, colour}, ...],
      page_index, note
    }
    key: list of uppercase A–D
    Writes xlsx; returns summary stats dict.
    """
    n = len(key)
    if n < 1:
        raise ValueError("empty answer key")

    by_class: dict[str, list] = defaultdict(list)
    for s in students:
        cc = s.get("class_code") or "UNKNOWN"
        by_class[cc].append(s)

    # stable class order: known codes then UNKNOWN
    def class_sort_key(c):
        if c == "UNKNOWN":
            return (1, c)
        return (0, c)

    classes = sorted(by_class.keys(), key=class_sort_key)
    # single sheet if one class (incl. all UNKNOWN)
    use_per_class = len(classes) > 1

    wb = Workbook()
    wb.remove(wb.active)
    thin = Side(style="thin", color="BBBBBB")
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    summary = {}

    sheet_specs = []
    if use_per_class:
        for c in classes:
            sheet_specs.append((c, by_class[c]))
    else:
        only = classes[0] if classes else "All"
        sheet_specs.append((only if only != "UNKNOWN" else "All", by_class.get(only, students)))

    for sheet_name, rows_in in sheet_specs:
        rows = sorted(
            rows_in,
            key=lambda r: (
                r.get("class_code") or "",
                r.get("class_no") if r.get("class_no") is not None else 999,
                r.get("page_index", 0),
            ),
        )
        safe = sheet_name[:31] or "Results"
        ws = wb.create_sheet(safe)

        ws["A1"] = f"{title} – MC results – {sheet_name}"
        ws["A1"].font = Font(bold=True, size=14)
        ws["A2"] = (
            "Green = correct (B/b OK); red = wrong / blank / double (score 0); "
            "yellow = OMR unclear — type A–D. CLSNP MC Marker."
        )
        ws["A2"].font = Font(italic=True, size=9, color="555555")

        H, K, F = 3, 4, 5
        L = F + len(rows) - 1 if rows else F
        heads = ["Class", "Class No.", "Name"] + [f"Q{q}" for q in range(1, n + 1)] + [f"Total (/{n})"]
        for j, h in enumerate(heads, 1):
            c = ws.cell(H, j, h)
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = FILL_HEADER
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = bd

        ws.cell(K, 3, "Correct answer")
        for q in range(n):
            ws.cell(K, 4 + q, key[q])
        ws.cell(K, 4 + n, n)
        for j in range(1, 5 + n):
            c = ws.cell(K, j)
            c.fill = FILL_KEY
            c.font = Font(bold=True)
            c.border = bd
            c.alignment = Alignment(horizontal="center")
        ws.cell(K, 3).alignment = Alignment(horizontal="left")

        totals = []
        for i, r in enumerate(rows):
            rr = F + i
            ws.cell(rr, 1, r.get("class_code") or "UNKNOWN")
            no = r.get("class_no")
            ws.cell(rr, 2, no if no is not None else "?")
            name = r.get("name") or ""
            ws.cell(rr, 3, name)
            answers = r.get("answers") or []
            score = 0
            for q in range(n):
                cell_info = answers[q] if q < len(answers) else {"ans": "?", "colour": "yellow"}
                ans = cell_info.get("ans", "?")
                colour = cell_info.get("colour", "yellow")
                cell = ws.cell(rr, 4 + q, ans)
                cell.fill = COLOUR_MAP.get(colour, FILL_YELLOW)
                cell.border = bd
                cell.alignment = Alignment(horizontal="center")
                if colour == "green":
                    score += 1
            tot_cell = ws.cell(rr, 4 + n, score)
            tot_cell.font = Font(bold=True)
            tot_cell.border = bd
            tot_cell.alignment = Alignment(horizontal="center")
            for j in range(1, 4):
                c = ws.cell(rr, j)
                c.border = bd
                c.alignment = Alignment(horizontal="left" if j == 3 else "center")
            totals.append(score)

        # empty class: still show key + note
        if not rows:
            ws.cell(F, 3, "(no students decoded)")
            for j in range(1, 5 + n):
                ws.cell(F, j).fill = FILL_YELLOW
                ws.cell(F, j).border = bd
            L = F

        P = L + 1
        A = L + 2
        ws.cell(P, 3, "% correct")
        for q in range(n):
            if rows:
                correct = sum(
                    1
                    for r in rows
                    if q < len(r.get("answers") or [])
                    and (r["answers"][q].get("colour") == "green")
                )
                pct = correct / len(rows)
            else:
                pct = 0.0
            c = ws.cell(P, 4 + q, pct)
            c.number_format = "0%"
        ws.cell(A, 3, f"Class average (/{n}) →")
        avg = (sum(totals) / len(totals)) if totals else 0.0
        c = ws.cell(A, 4 + n, avg)
        c.number_format = "0.0"
        for rr in (P, A):
            for j in range(1, 5 + n):
                cell = ws.cell(rr, j)
                cell.fill = FILL_SUM
                cell.font = Font(bold=True)
                cell.border = bd
                cell.alignment = Alignment(horizontal="center")
            ws.cell(rr, 3).alignment = Alignment(horizontal="left")

        ws.column_dimensions["A"].width = 8
        ws.column_dimensions["B"].width = 11
        ws.column_dimensions["C"].width = 18
        for q in range(n):
            ws.column_dimensions[CL(4 + q)].width = 5.5
        ws.column_dimensions[CL(4 + n)].width = 11
        ws.row_dimensions[H].height = 20
        ws.freeze_panes = f"D{F}"
        ws.sheet_view.zoomScale = 90

        summary[sheet_name] = {
            "n": len(rows),
            "mean": avg,
            "totals": totals,
        }

    wb.save(out_path)
    return summary
